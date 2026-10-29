import copy
import importlib.util
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile
from PIL import Image, ImageDraw
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SKILL = ROOT / "skills/scan-pdf-to-word"
sys.path.insert(0, str(SKILL / "scripts"))
from build_docx import Builder, validate
from check_docx import check
from common import digest
from pdf_pages import render, page_selection
from bootstrap import incompatible_packages


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.work = Path(self.temp.name) / "中文 路径"
        self.work.mkdir()
        self.data = json.loads((ROOT / "examples/demo-transcript.json").read_text(encoding="utf-8"))

    def tearDown(self):
        self.temp.cleanup()

    def test_native_text_tables_comments_and_editability(self):
        path = self.work / "可编辑 文档.docx"
        report = Builder(self.data, self.work).build(path)
        original_hash = digest(path)
        result = check(path, self.data)
        self.assertTrue(result["automated_checks_passed"], result)
        self.assertEqual(original_hash, digest(path), "Checking must not mutate the input")
        self.assertEqual(report["native_tables"], 2)
        self.assertEqual(report["images"], 0)
        self.assertEqual(report["comments"], 1)
        document = Document(path)
        self.assertIn("繁體字", "".join(p.text for p in document.paragraphs))
        self.assertEqual(document.tables[0].cell(1, 1).text, "")
        self.assertIn("合併標題", document.tables[1].cell(0, 2).text)
        self.assertEqual(document.tables[1].cell(0, 0)._tc, document.tables[1].cell(0, 2)._tc)
        widths = [int(c._tc.tcPr.tcW.get(qn("w:w"))) for c in document.tables[0].rows[1].cells]
        self.assertGreater(widths[1], 3 * widths[0])

    def test_missing_text_is_detected(self):
        path = self.work / "test.docx"
        Builder(self.data, self.work).build(path)
        altered = copy.deepcopy(self.data)
        altered["pages"][0]["blocks"][0]["text"] += "漏字"
        result = check(path, altered)
        self.assertFalse(result["automated_checks_passed"])
        self.assertTrue(any("source page 1" in x for x in result["failures"]))

    def test_no_overwrite(self):
        path = self.work / "test.docx"
        Builder(self.data, self.work).build(path)
        original_hash = digest(path)
        with self.assertRaises(ValueError):
            Builder(self.data, self.work).build(path)
        self.assertEqual(original_hash, digest(path))

    def test_bad_merge_and_unordered_pages_rejected(self):
        altered = copy.deepcopy(self.data)
        altered["pages"][1]["blocks"][1]["rows"][0][1] = "would be lost"
        with self.assertRaises(ValueError):
            validate(altered)
        altered = copy.deepcopy(self.data)
        altered["pages"].reverse()
        with self.assertRaises(ValueError):
            validate(altered)

    def test_scan_pdf_rendering_and_page_identity(self):
        images = []
        for i in range(2):
            image = Image.new("RGB", (600, 800), "white")
            ImageDraw.Draw(image).text((50, 50), f"SCAN PAGE {i + 1}", fill="black")
            images.append(image)
        source = self.work / "原件 scan.pdf"
        images[0].save(source, save_all=True, append_images=images[1:])
        old_hash = digest(source)
        out = self.work / "source pages"
        manifest = render(source, out, dpi=100, selected="2")
        self.assertEqual(manifest["source_page_count"], 2)
        self.assertEqual(manifest["pages"][0]["source_page"], 2)
        self.assertEqual(manifest["pages"][0]["text_layer_characters"], 0)
        self.assertTrue((out / "page-0002.png").is_file())
        self.assertEqual(digest(source), old_hash)
        with self.assertRaises(ValueError):
            render(source, out, dpi=100)
        with self.assertRaises(ValueError):
            page_selection("3", 2)

    def test_skill_references_and_metadata(self):
        text = (SKILL / "SKILL.md").read_text(encoding="utf-8")
        self.assertTrue(text.startswith("---\nname: scan-pdf-to-word\n"))
        for rel in re.findall(r"\]\((references/[^)]+)\)", text):
            self.assertTrue((SKILL / rel).is_file(), rel)
        self.assertNotIn("/Users/", text)
        self.assertNotIn("TODO", text)

    def test_old_word_library_triggers_private_environment(self):
        versions = {"python-docx": "0.8.11", "pypdfium2": "5.14.0", "Pillow": "12.3.0"}
        with patch("bootstrap.importlib.metadata.version", side_effect=versions.__getitem__):
            self.assertEqual(incompatible_packages(), ["python-docx"])

    def test_installer_and_refusal_to_overwrite(self):
        fake_home = self.work / "installer home"
        if sys.platform == "win32":
            command = ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File",
                       str(ROOT / "install.ps1"), "-Target", "all", "-UserRoot", str(fake_home)]
        else:
            command = ["sh", str(ROOT / "install.sh"), "all", str(fake_home)]
        subprocess.run(command, check=True, capture_output=True)
        for parent in (".agents", ".claude"):
            installed = fake_home / parent / "skills/scan-pdf-to-word"
            self.assertEqual(digest(installed / "SKILL.md"), digest(SKILL / "SKILL.md"))
            self.assertTrue((installed / "scripts/build_docx.py").exists())
        self.assertNotEqual(subprocess.run(command, capture_output=True).returncode, 0)

    def test_release_package_contains_skill_and_no_case_documents(self):
        subprocess.run([sys.executable, str(ROOT / "package_release.py")], check=True, capture_output=True)
        with zipfile.ZipFile(ROOT / "dist/scan-pdf-to-word.zip") as z:
            names = z.namelist()
            self.assertIn("scan-pdf-to-word/SKILL.md", names)
            self.assertIn("scan-pdf-to-word/LICENSE", names)
            self.assertIn("scan-pdf-to-word/scripts/vision.swift", names)
            self.assertFalse(any(x.endswith((".pdf", ".docx", ".pyc", ".traineddata")) for x in names))
        with zipfile.ZipFile(ROOT / "dist/scan-pdf-to-word-starter.zip") as z:
            self.assertIn("scan-pdf-to-word-skill/install.ps1", z.namelist())
            self.assertFalse(any("__pycache__" in x for x in z.namelist()))


if __name__ == "__main__":
    unittest.main()
