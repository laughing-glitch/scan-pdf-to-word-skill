"""Check native Word content against the transcript; this does not prove OCR accuracy."""
import argparse
from pathlib import Path
import re
import sys
import tempfile
import zipfile
from lxml import etree
from docx import Document
from build_docx import expected_page_text, validate
from common import digest, emit, read_json, write_json

NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def compact(text):
    return re.sub(r"\s+", "", text)


def check(path, schema, rendered_pdf=None):
    schema = validate(schema)
    with zipfile.ZipFile(path) as z:
        if z.testzip():
            raise ValueError("Corrupt DOCX archive")
        root = etree.fromstring(z.read("word/document.xml"))
        comments = etree.fromstring(z.read("word/comments.xml")) if "word/comments.xml" in z.namelist() else None
    pages = [""]
    for child in root.find("w:body", NS):
        if child.xpath('.//w:br[@w:type="page"]', namespaces=NS):
            pages.append("")
        pages[-1] += "".join(child.xpath(".//w:t/text()", namespaces=NS))
    failures = []
    if len(pages) != len(schema["pages"]):
        failures.append("Hard page-break count differs from transcript")
    for i, page in enumerate(schema["pages"]):
        if i >= len(pages) or compact(expected_page_text(page)) != compact(pages[i]):
            failures.append(f"Native text differs from transcript on source page {page['source_page']}")
    ids = comments.xpath("./w:comment/@w:id", namespaces=NS) if comments is not None else []
    for tag in ("commentRangeStart", "commentRangeEnd", "commentReference"):
        if sorted(root.xpath(f".//w:{tag}/@w:id", namespaces=NS)) != sorted(ids):
            failures.append("Comment anchor mismatch: " + tag)
    textboxes = len(root.xpath(".//w:txbxContent", namespaces=NS))
    tables = len(root.xpath(".//w:tbl", namespaces=NS))
    if not any(compact(x) for x in pages):
        failures.append("No native body text found")
    d = Document(path)
    p = next((p for p in d.paragraphs if p.text), d.add_paragraph())
    p.add_run("EDIT_SAVE_REOPEN_PROBE")
    if d.tables:
        d.tables[0].cell(0, 0).paragraphs[0].add_run("CELL_EDIT_PROBE")
    with tempfile.TemporaryDirectory() as temp:
        copy = Path(temp) / "editable-check.docx"
        d.save(copy)
        reopened = Document(copy)
        if not any("EDIT_SAVE_REOPEN_PROBE" in p.text for p in reopened.paragraphs):
            failures.append("Paragraph edit/save/reopen failed")
        if d.tables and "CELL_EDIT_PROBE" not in reopened.tables[0].cell(0, 0).text:
            failures.append("Table edit/save/reopen failed")
    rendered = None
    if rendered_pdf:
        import pypdfium2 as pdfium
        with pdfium.PdfDocument(str(rendered_pdf)) as pdf:
            rendered = len(pdf)
        if schema.get("document", {}).get("preserve_page_count", True) and rendered != len(schema["pages"]):
            failures.append("Rendered page count differs from source-page transcript")
    return {"docx_sha256": digest(path), "source_pages_in_transcript": len(schema["pages"]),
            "rendered_pages": rendered, "native_tables": tables, "textboxes": textboxes,
            "images": len(d.inline_shapes), "comments": len(ids), "failures": failures,
            "automated_checks_passed": not failures,
            "visual_source_proofreading": "not verified by this script",
            "visual_render_review": "not verified by this script",
            "accuracy_percentage": None}


if __name__ == "__main__":
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("docx", type=Path)
    p.add_argument("--transcript", required=True, type=Path)
    p.add_argument("--rendered-pdf", type=Path)
    p.add_argument("--report", type=Path)
    a = p.parse_args()
    result = check(a.docx, read_json(a.transcript), a.rendered_pdf)
    if a.report:
        if a.report.exists():
            raise SystemExit("Report already exists; choose a new path")
        write_json(a.report, result)
    emit(result)
    sys.exit(0 if result["automated_checks_passed"] else 1)
