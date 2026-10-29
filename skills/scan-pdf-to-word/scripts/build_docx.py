"""Build native editable Word elements from a visually proofread JSON transcript."""
import argparse
import io
import math
from pathlib import Path
import sys
from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image
from common import emit, read_json

ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
         "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}


def plain(value):
    if isinstance(value, str):
        return value
    if isinstance(value, list) and all(isinstance(x, dict) and isinstance(x.get("text"), str) for x in value):
        return "".join(x["text"] for x in value)
    raise ValueError("Text must be a string or a list of run objects with text fields")


def paragraphs(cell):
    if isinstance(cell, dict):
        return cell.get("paragraphs", [{"text": cell.get("text", "")}])
    return [{"text": cell}]


def expected_page_text(page):
    result = []
    for block in page["blocks"]:
        if block["type"] == "paragraph":
            result.append(plain(block.get("text", "")))
        elif block["type"] == "table":
            for row in block["rows"]:
                for cell in row:
                    result.extend(plain(p.get("text", "")) for p in paragraphs(cell))
        elif block["type"] == "image" and block.get("caption"):
            result.append(plain(block["caption"]))
    return "\n".join(result)


def validate(data):
    if data.get("schema_version") != 1 or not isinstance(data.get("pages"), list) or not data["pages"]:
        raise ValueError("Expected schema_version: 1 and non-empty pages")
    numbers = []
    for page in data["pages"]:
        number = page.get("source_page")
        if not isinstance(number, int) or number < 1 or number in numbers:
            raise ValueError("source_page must be a unique positive integer")
        numbers.append(number)
        if not page.get("blocks"):
            raise ValueError(f"Source page {number} has no blocks; represent intentional blanks explicitly")
        for block in page["blocks"]:
            kind = block.get("type")
            if kind not in ("paragraph", "table", "image"):
                raise ValueError(f"Unknown block type: {kind}")
            if kind == "paragraph":
                plain(block.get("text", ""))
            elif kind == "table":
                widths, rows = block["widths_cm"], block["rows"]
                if not widths or any(not isinstance(x, (int, float)) or not math.isfinite(x) or x <= 0 for x in widths):
                    raise ValueError("Table widths must be positive finite numbers")
                if not rows or any(len(row) != len(widths) for row in rows):
                    raise ValueError("Every table row must match widths_cm")
                covered = set()
                for r0, c0, r1, c1 in block.get("merges", []):
                    if not (0 <= r0 <= r1 < len(rows) and 0 <= c0 <= c1 < len(widths)):
                        raise ValueError("Table merge is out of bounds")
                    area = {(r, c) for r in range(r0, r1 + 1) for c in range(c0, c1 + 1)}
                    if covered & area:
                        raise ValueError("Overlapping table merges")
                    covered |= area
                    for r, c in area - {(r0, c0)}:
                        if any(plain(p.get("text", "")).strip() for p in paragraphs(rows[r][c])):
                            raise ValueError("Merged cells must store all content in the top-left cell")
                for row in rows:
                    for cell in row:
                        for p in paragraphs(cell):
                            plain(p.get("text", ""))
            elif block.get("purpose") not in ("illustration", "handwriting", "diagram"):
                raise ValueError("Images require purpose: illustration, handwriting or diagram; don't embed full pages as editable text")
    if numbers != sorted(numbers):
        raise ValueError("Source pages must remain in order")
    return data


class Builder:
    def __init__(self, data, base):
        self.data, self.base = validate(data), Path(base)
        self.doc = Document()
        self.source_page = None
        self.settings = data.get("document", {})
        default_font = "Songti TC" if sys.platform == "darwin" else "SimSun" if sys.platform == "win32" else "Noto Serif CJK TC"
        self.cjk = self.settings.get("font_east_asia", default_font)
        self.size = self.settings.get("font_size_pt", 12)
        sec = self.doc.sections[0]
        size = self.settings.get("page_size_cm", [21, 29.7])
        sec.page_width, sec.page_height = map(Cm, size)
        margins = self.settings.get("margins_cm", [1.8, 1.8, 1.8, 1.8])
        sec.top_margin, sec.right_margin, sec.bottom_margin, sec.left_margin = map(Cm, margins)
        sec.header_distance = sec.footer_distance = Cm(0.7)
        self.doc.core_properties.title = self.settings.get("title", "")
        self.doc.core_properties.author = ""
        self.doc.core_properties.last_modified_by = ""
        for name in ("Normal", "Title", "Heading 1", "Heading 2", "Header", "Footer"):
            style = self.doc.styles[name]
            style.font.name = "Times New Roman"
            style.font.size = Pt(self.size)
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), self.cjk)
            style.paragraph_format.space_after = Pt(6)
        for border in list(self.doc.styles.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
        for area in ("header", "footer"):
            if self.settings.get(area):
                self.add_runs(getattr(sec, area).paragraphs[0], self.settings[area], 8)

    def add_runs(self, p, value, size=None, bold=False):
        plain(value)
        runs = [{"text": value}] if isinstance(value, str) else value
        for item in runs:
            r = p.add_run(item["text"])
            r.font.name = "Times New Roman"
            r.font.size = Pt(item.get("size_pt", size or self.size))
            r.bold, r.italic, r.underline = item.get("bold", bold), item.get("italic", False), item.get("underline", False)
            pr = r._element.get_or_add_rPr()
            pr.get_or_add_rFonts().set(qn("w:eastAsia"), self.cjk)
            lang = OxmlElement("w:lang")
            lang.set(qn("w:eastAsia"), self.settings.get("language", "zh-TW"))
            pr.append(lang)
            if item.get("comment"):
                if not item["text"]:
                    raise ValueError("Comment anchors must contain text")
                note = f"原 PDF 第 {self.source_page} 頁：{item['comment']}"
                self.doc.add_comment(r, note, author="轉錄校對", initials="校")

    def paragraph(self, container, block, existing=None):
        p = existing if existing is not None else container.add_paragraph()
        if block.get("style"):
            p.style = block["style"]
        pf = p.paragraph_format
        pf.space_after = Pt(block.get("space_after_pt", 7))
        pf.line_spacing = Pt(block.get("line_spacing_pt", block.get("size_pt", self.size) * 1.55))
        pf.keep_with_next = block.get("keep_with_next", False)
        pf.keep_together = block.get("keep_together", True)
        pf.widow_control = True
        pf.left_indent = Cm(block.get("left_indent_cm", 0))
        pf.first_line_indent = Cm(block.get("first_line_indent_cm", 0))
        p.alignment = ALIGN[block.get("align", "left")]
        self.add_runs(p, block.get("text", ""), block.get("size_pt"), block.get("bold", False))
        return p

    def table(self, block):
        t = self.doc.add_table(rows=len(block["rows"]), cols=len(block["widths_cm"]))
        t.alignment, t.autofit = WD_TABLE_ALIGNMENT.CENTER, False
        for i, width in enumerate(block["widths_cm"]):
            t.columns[i].width = Cm(width)
        for row in t.rows:
            for ci, cell in enumerate(row.cells):
                cell.width = Cm(block["widths_cm"][ci])
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement("w:" + side)
            el.set(qn("w:val"), block.get("border", "single"))
            el.set(qn("w:sz"), "6")
            el.set(qn("w:color"), "808080")
            borders.append(el)
        t._tbl.tblPr.append(borders)
        for r0, c0, r1, c1 in block.get("merges", []):
            t.cell(r0, c0).merge(t.cell(r1, c1))
        done = set()
        for ri, row in enumerate(t.rows):
            row.height = Cm(block.get("min_row_heights_cm", [])[ri] if ri < len(block.get("min_row_heights_cm", [])) else 0.7)
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            if ri < block.get("header_rows", 0):
                row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for ci, cell in enumerate(row.cells):
                if cell._tc in done:
                    continue
                done.add(cell._tc)
                cell.text = ""
                content = block["rows"][ri][ci]
                for pi, item in enumerate(paragraphs(content)):
                    item = {"size_pt": block.get("size_pt", self.size), "space_after_pt": 3, **item}
                    self.paragraph(cell, item, cell.paragraphs[0] if pi == 0 else None)
        self.paragraph(self.doc, {"text": "", "size_pt": 1, "line_spacing_pt": 1, "space_after_pt": 3})

    def image(self, block):
        path = self.base / block["path"]
        with Image.open(path) as original:
            image = original.copy()
        if block.get("crop"):
            x0, y0, x1, y1 = block["crop"]
            if not (0 <= x0 < x1 <= 1 and 0 <= y0 < y1 <= 1):
                raise ValueError("Image crop must be normalized x0,y0,x1,y1")
            image = image.crop((round(x0 * image.width), round(y0 * image.height), round(x1 * image.width), round(y1 * image.height)))
        buf = io.BytesIO()
        image.save(buf, format="PNG")
        buf.seek(0)
        p = self.doc.add_paragraph()
        p.alignment = ALIGN[block.get("align", "center")]
        picture = p.add_run().add_picture(buf, width=Cm(block.get("width_cm", 6)))
        picture._inline.docPr.set("descr", block.get("alt", block["purpose"]))
        if block.get("caption"):
            self.paragraph(self.doc, {"text": block["caption"], "size_pt": 10, "align": "center"})

    def build(self, output):
        output = Path(output)
        if output.exists():
            raise ValueError(f"Refusing to overwrite existing file: {output}")
        for i, page in enumerate(self.data["pages"]):
            self.source_page = page["source_page"]
            if i:
                p = self.paragraph(self.doc, {"text": "", "size_pt": 1, "line_spacing_pt": 1, "space_after_pt": 0})
                p.add_run().add_break(WD_BREAK.PAGE)
            for block in page["blocks"]:
                if block["type"] == "paragraph":
                    self.paragraph(self.doc, block)
                else:
                    getattr(self, block["type"])(block)
        output.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(output)
        return {"output": str(output.resolve()), "source_pages": len(self.data["pages"]),
                "native_tables": len(self.doc.tables), "images": len(self.doc.inline_shapes),
                "comments": len(self.doc.comments), "visual_review": "required"}


if __name__ == "__main__":
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("transcript", type=Path)
    p.add_argument("--out", required=True, type=Path)
    a = p.parse_args()
    emit(Builder(read_json(a.transcript), a.transcript.resolve().parent).build(a.out))
